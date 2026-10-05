"""Spreadsheet export: what a coordinator sees, in CSV or Excel."""

import csv
import io

from fastapi.testclient import TestClient

from fairtriage.api import app
from fairtriage import service
from fairtriage.schemas import LodgeIn

c = TestClient(app)


def seed():
    service.lodge(LodgeIn(text="the power point is sparking", community="Wadeye",
                          address="Lot 12", phone="0400 000 000"))
    service.lodge(LodgeIn(text="the kitchen cupboard door came off the hinge",
                          community="Darwin (Parap)", address="4 Smith Street"))
    service.lodge(LodgeIn(text="its fine now dont worry about it", community="Galiwinku"))


def test_csv_is_the_ranked_queue():
    seed()
    r = c.get("/api/export?format=csv")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert [x["Priority"] for x in rows] == ["Immediate", "Routine"]      # ranked order
    assert rows[0]["Address"] == "Lot 12" and rows[0]["Position"] == "1"
    assert rows[0]["Tenant's words"]


def test_all_requests_includes_those_not_in_the_queue():
    seed()
    rows = list(csv.DictReader(io.StringIO(
        c.get("/api/export?format=csv&scope=all").content.decode("utf-8-sig"))))
    assert len(rows) == 3
    assert any(x["Status"] == "awaiting confirmation" for x in rows)


def test_filters_apply():
    seed()
    rows = list(csv.DictReader(io.StringIO(
        c.get("/api/export?format=csv&remote=true").content.decode("utf-8-sig"))))
    assert [x["Community"] for x in rows] == ["Wadeye"]


def test_area_filter_and_column():
    seed()
    rows = list(csv.DictReader(io.StringIO(
        c.get("/api/export?format=csv&area=Top%20End").content.decode("utf-8-sig"))))
    assert [(x["Area"], x["Community"]) for x in rows] == [("Top End", "Wadeye")]
    assert {r["area"] for r in c.get("/api/queue").json()} == {"Top End", "Darwin"}


def test_xlsx_opens_with_headers_and_rows():
    from openpyxl import load_workbook
    seed()
    r = c.get("/api/export?format=xlsx")
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws["B1"].value == "Reference" and ws.max_row == 3
    assert ws.auto_filter.ref and ws.freeze_panes == "C2"


def test_bad_format_is_refused():
    assert c.get("/api/export?format=pdf").status_code == 422


def test_tenant_text_never_becomes_a_spreadsheet_formula():
    """Regression: '=HYPERLINK(...)' typed by a tenant was exported as a live
    formula (CSV and Excel). It must arrive as plain text."""
    import io
    from openpyxl import load_workbook
    from fairtriage import export, service
    from fairtriage.schemas import LodgeIn
    service.lodge(LodgeIn(text='=HYPERLINK("http://evil","x") the tap drips', community="Darwin (Parap)"))
    data = export.rows("queue", None, None, None, None)
    csv = export.to_csv(data).decode("utf-8")
    assert '"\'=HYPERLINK' in csv or ",'=HYPERLINK" in csv
    ws = load_workbook(io.BytesIO(export.to_xlsx(data))).active
    cells = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and "HYPERLINK" in c.value]
    assert cells and all(c.startswith("'") for c in cells)
    assert all(c.data_type != "f" for row in ws.iter_rows() for c in row)



def test_download_filters_by_several_priorities_and_past_target():
    """The queue's priority filter: any mix of tiers, and past target only."""
    from datetime import datetime, timedelta, timezone
    from fairtriage import export, service
    from fairtriage.schemas import LodgeIn

    def ago(d):
        return (datetime.now(timezone.utc) - timedelta(days=d)).isoformat(timespec="seconds")

    service.lodge(LodgeIn(text="the power point is sparking", community="Darwin (Parap)"))
    service.lodge(LodgeIn(text="the only toilet is blocked and will not flush", community="Darwin (Parap)"),
                  lodged_at=ago(9))                                   # Urgent, past its 2-day target
    service.lodge(LodgeIn(text="the kitchen tap is dripping", community="Darwin (Parap)"))
    both = export.rows("queue", "Immediate,Urgent")
    assert sorted(r["tier"] for r in both) == ["Immediate", "Urgent"]
    late = export.rows("queue", None, past=True)
    assert [r["tier"] for r in late] == ["Urgent"] and all(r["past_target"] for r in late)
    assert len(export.rows("queue")) == 3
