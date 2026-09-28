"""Spreadsheet export of repair requests: CSV or Excel.

Built from the same views the screens use, so a download shows exactly what a
coordinator sees: the ranked queue (optionally filtered) or every request.
"""

from __future__ import annotations

import csv
import io
from datetime import datetime, timezone

from .service import all_requests_view, queue_view

COLUMNS = [
    # (header, key, width)
    ("Position", "position", 9), ("Reference", "request_id", 18), ("Priority", "tier", 11),
    ("Need score", "need", 10), ("Status", "status", 16), ("Reported (UTC)", "lodged_at", 20),
    ("Days waiting", "days_open", 12), ("Target", "target_label", 16),
    ("% of target", "pct_of_target", 11), ("Past target", "past_target", 11),
    ("Area", "area", 14), ("Community", "community", 22), ("Address", "address", 30), ("Phone", "phone", 15),
    ("Remote", "remote", 8), ("Trade", "trade", 14), ("Tenant's words", "evidence", 60),
    ("Expected wait from (days)", "wait_low", 14), ("Expected wait to (days)", "wait_high", 14),
    ("Flags", "flags", 40), ("Trip", "trip_id", 14), ("Expected arrival (UTC)", "eta_at", 20),
]

TIER_FILL = {"Immediate": "FDECEA", "Urgent": "FDF3E2", "Routine": "E8F3ED"}


def rows(scope: str = "queue", tier: str | None = None, remote: bool | None = None,
         q: str | None = None, area: str | None = None) -> list[dict]:
    data = queue_view(tier) if scope == "queue" else all_requests_view()
    if scope != "queue" and tier:
        data = [r for r in data if r.get("tier") == tier]
    if area:
        data = [r for r in data if r.get("area") == area]
    if remote is not None:
        data = [r for r in data if r.get("remote") == remote]
    if q:
        needle = q.lower()
        data = [r for r in data if any(needle in str(r.get(k) or "").lower()
                                       for k in ("request_id", "community", "address", "text", "trade"))]
    return data


def _cell(r: dict, key: str):
    v = r.get(key)
    if key == "evidence":
        v = v or r.get("text", "")
    if key == "flags":
        return "; ".join(f["code"].replace("_", " ") for f in (v or []))
    if key in ("remote", "past_target"):
        return "" if v is None else ("yes" if v else "no")
    if key == "status" and v:
        return v.replace("_", " ")
    return "" if v is None else v


def filename(fmt: str) -> str:
    return f"fairtriage-requests-{datetime.now(timezone.utc):%Y-%m-%d}.{fmt}"


def to_csv(data: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([h for h, _, _ in COLUMNS])
    for r in data:
        w.writerow([_cell(r, k) for _, k, _ in COLUMNS])
    # BOM so Excel opens it as UTF-8 (tenants' words carry curly quotes)
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(data: list[dict], title: str = "Requests") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([h for h, _, _ in COLUMNS])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2B3A8F")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    tier_col = [k for _, k, _ in COLUMNS].index("tier") + 1
    for r in data:
        ws.append([_cell(r, k) for _, k, _ in COLUMNS])
        fill = TIER_FILL.get(r.get("tier", ""))
        if fill:
            ws.cell(row=ws.max_row, column=tier_col).fill = PatternFill("solid", fgColor=fill)
    for i, (_, _, width) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
